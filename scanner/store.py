"""The running list. data/master.json is the scanner's memory; the Google Sheet is the human view.

Rules: new finds append, nothing is deleted, duplicates merge (same canonical link or
same company+role key), and your edits in the Sheet win for the columns you own and
for APPLIED / RULED OUT placement.
"""
import json
from datetime import date
from pathlib import Path

from .classify import enrich, parse_date, tab_for
from .model import COLUMNS, HEADERS, TABS, USER_OWNED, canon_url

PATH = Path(__file__).resolve().parent.parent / "data/master.json"
LOCKED_TABS = {"APPLIED", "RULED OUT"}


def _term(role):
    import re
    m = re.search(r"\b(spring|summer|fall|autumn|winter)\b[^0-9]{0,12}(20\d\d)?", role or "", re.I)
    return (m.group(1).lower(), m.group(2)) if m else None


class Store:
    def __init__(self, path=PATH):
        self.path = path
        d = json.load(open(path)) if path.exists() else {}
        self.rows = d.get("rows", {})          # key -> row dict (+ "tab")
        self.urls = d.get("urls", {})          # canonical url -> key
        self.rubric = set(d.get("rubric_scored", []))
        self.dirty = set(d.get("dirty", []))     # changed by scanner.update, not yet pushed

    def save(self):
        tmp = self.path.with_suffix(".tmp")
        json.dump({"rows": self.rows, "urls": self.urls, "rubric_scored": sorted(self.rubric),
                   "dirty": sorted(self.dirty)},
                  open(tmp, "w"), indent=1, default=str)
        tmp.replace(self.path)

    def find(self, row):
        u = canon_url(row.get("link"))
        if u and u in self.urls and self.urls[u] in self.rows:
            return self.urls[u]
        k = row["key"]
        if k in self.rows:
            cu = canon_url(self.rows[k].get("link"))
            if u and cu and u != cu and _term(row["role"]) != _term(self.rows[k]["role"]):
                # Spring and Summer of the same program are two postings: keep both. Same title and term
                # under a second requisition is a repost of one role, so it merges.
                import hashlib
                row["key"] = f"{k}#{hashlib.sha1(u.encode()).hexdigest()[:6]}"
                return row["key"] if row["key"] in self.rows else None
            return k
        return None

    def _index(self, key, row):
        u = canon_url(row.get("link"))
        if u:
            self.urls.setdefault(u, key)

    def add(self, row, today, tab=None, why=None):
        """Merge one row. Returns 'new', 'merged' or 'dup'."""
        if why is None:
            why = enrich(row)
        key = self.find(row)
        if key:
            cur = self.rows[key]
            changed = upgraded = False
            for c in COLUMNS:
                if c in ("key", "tab") or not row.get(c):
                    continue
                if c == "source" and row[c] not in cur["source"].split("; "):
                    cur["source"] = f"{cur['source']}; {row[c]}".strip("; ")
                    changed = True
                elif c == "verified" and row[c] == "y" and cur.get(c) != "y":
                    # a firm-page read beats whatever the earlier source said about class year and dates
                    cur[c], changed = "y", True
                    for f in ("class_year", "deadline", "opens"):
                        if row.get(f):
                            cur[f] = row[f]
                    upgraded = True
                elif not cur.get(c) and c not in USER_OWNED:
                    cur[c], changed = row[c], True
            self._index(key, row)
            if upgraded and cur["tab"] not in LOCKED_TABS and key not in self.rubric:
                cur["fit"] = ""
                why2 = enrich(cur)
                if why2:
                    cur["tab"] = "RULED OUT"
                    cur["notes"] = f"RULED OUT: {why2}. {cur.get('notes', '')}".strip()
            return "merged" if changed else "dup"
        key = row["key"]
        row["date_found"] = row.get("date_found") or today.isoformat()
        row["status"] = row.get("status") or "new"
        row["verified"] = row.get("verified") or "n"
        if why:
            row["tab"] = "RULED OUT"
            row["notes"] = f"RULED OUT: {why}. {row.get('notes', '')}".strip()
        else:
            row["tab"] = tab or tab_for(row, today, key in self.rubric)
        self.rows[key] = row
        self._index(key, row)
        return "new"

    def retab(self, today):
        """Promote rows whose deadline came inside 30 days; retire unapplied rows whose deadline passed."""
        moves = []
        for k, r in self.rows.items():
            if r["tab"] in LOCKED_TABS:
                continue
            d = parse_date(r.get("deadline"))
            if d and d < today and r.get("status") not in ("applied", "staged"):
                r["tab"] = "RULED OUT"
                r["notes"] = f"RULED OUT: deadline {d} passed. {r.get('notes', '')}".strip()
                moves.append((k, "RULED OUT"))
                continue
            want = tab_for(r, today, k in self.rubric)
            if want != r["tab"]:
                r["tab"] = want
                moves.append((k, want))
        return moves

    def pull_sheet(self, sheet_rows):
        """Apply your edits from the Sheet. sheet_rows: list of row dicts with 'tab'."""
        n = 0
        for s in sheet_rows:
            k = s.get("key")
            if not k:
                continue
            if k not in self.rows:  # a row added by hand
                self.rows[k] = {c: s.get(c, "") for c in COLUMNS} | {"tab": s["tab"]}
                n += 1
                continue
            cur = self.rows[k]
            for c in USER_OWNED | {"fit"}:
                if s.get(c, "") != "" and s.get(c) != cur.get(c):
                    cur[c] = s[c]
                    n += 1
            if s["tab"] != cur["tab"] and (s["tab"] in LOCKED_TABS or cur["tab"] in LOCKED_TABS):
                cur["tab"] = s["tab"]
                n += 1
        return n

    def by_tab(self, tab):
        rows = [r for r in self.rows.values() if r["tab"] == tab]
        return sorted(rows, key=lambda r: (parse_date(r.get("deadline")) or date.max, -float(r.get("fit") or 0)))

    def export_xlsx(self, path):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        wb.remove(wb.active)
        for tab in TABS:
            ws = wb.create_sheet(tab)
            ws.append(HEADERS)
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="1F3864" if tab != "URGENT" else "C00000")
            for r in self.by_tab(tab):
                ws.append([r.get(c, "") for c in COLUMNS])
            ws.freeze_panes = "C2"
            for col, w in zip("ABCDEFGHIJKLMNOPQR", [22, 44, 6, 15, 12, 12, 22, 40, 22, 9, 8, 10, 11, 20, 24, 14, 50, 10]):
                ws.column_dimensions[col].width = w
            ws.column_dimensions["R"].hidden = True
        wb.save(path)
